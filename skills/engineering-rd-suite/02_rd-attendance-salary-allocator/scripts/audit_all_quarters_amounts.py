import os
import xlrd
import openpyxl
from pathlib import Path

sum_dir = Path(r"C:\JunzanLi_project\课题申报横向\人工费考勤表与分摊表\季度汇总表")
res_dir = Path(r"C:\JunzanLi_project\课题申报横向\人工费考勤表与分摊表\结果\Antigravity_批量生成结果_21-B-02")

sum_files = sorted(sum_dir.glob("*.xls"))
proj_keywords = ["21-B-02", "长距离轨行区", "长距离"]

print("="*95)
print(f"{'季度文件':38s} | {'汇总表法定总额':14s} | {'3个月分摊合计':14s} | {'差额(元)':10s} | {'状态'}")
print("="*95)

grand_target_all = 0.0
grand_alloc_all = 0.0

q_to_months = {1: "1-3月", 2: "4-6月", 3: "7-9月", 4: "10-12月"}

for sf in sum_files:
    fname = sf.name
    import re
    m_match = re.search(r'(\d{4})年.*([1-4])季度', fname)
    if not m_match:
        continue
    year = int(m_match.group(1))
    q = int(m_match.group(2))
    m_str = q_to_months[q]
    
    # Find matching file in res_dir
    res_candidates = list(res_dir.glob(f"{year}年{m_str}*21-B-02*.xlsx"))
    if not res_candidates:
        continue
        
    target_res = None
    for c in res_candidates:
        if "最新重算版" in c.name:
            target_res = c
            break
    if not target_res:
        target_res = res_candidates[0]
        
    # Read summary target
    wb_s = xlrd.open_workbook(sf)
    snames = wb_s.sheet_names()
    entry_sname = None
    for s in snames:
        if any(k in s for k in ['分配表入账', '职工薪酬分配明细', '分配明细', '职工薪酬']):
            if s != '汇总表':
                entry_sname = s
                break
    if not entry_sname and len(snames) > 1:
        entry_sname = snames[1]
        
    es = wb_s.sheet_by_name(entry_sname)
    proj_rate_col = None
    for c in range(es.ncols):
        txt = " ".join([str(es.cell_value(r, c)) for r in range(min(5, es.nrows))])
        if any(k in txt for k in proj_keywords):
            proj_rate_col = c
            break
            
    proj_fee_col = proj_rate_col + 1 if proj_rate_col is not None else None
    name_col = 1
    for c in range(min(5, es.ncols)):
        for r in range(min(5, es.nrows)):
            if str(es.cell_value(r, c)).strip() == '姓名':
                name_col = c
                break
                
    target_fee_sum = 0.0
    people_targets = {}
    for r in range(es.nrows):
        v_name = str(es.cell_value(r, name_col)).strip()
        val_fee = es.cell_value(r, proj_fee_col) if proj_fee_col < es.ncols else 0.0
        val_rate = es.cell_value(r, proj_rate_col)
        if v_name and v_name not in ['姓名', '合计', '小计', ''] and val_rate not in ['', None, 0, 0.0]:
            try:
                f_num = float(val_fee) if val_fee not in ['', None] else 0.0
                people_targets[v_name] = f_num
            except ValueError:
                pass
        elif v_name == '合计' and val_fee not in ['', None]:
            try:
                target_fee_sum = float(val_fee)
            except ValueError:
                pass
                
    if target_fee_sum == 0.0 and people_targets:
        target_fee_sum = sum(people_targets.values())

    wb_r = openpyxl.load_workbook(target_res, data_only=False)
    monthly_sheet_names = [s for s in wb_r.sheetnames if '考勤' not in s]
    
    alloc_sum = 0.0
    for ms_name in monthly_sheet_names:
        ws_p = wb_r[ms_name]
        ws_a = wb_r[f"{ms_name}考勤表"]
        
        for r in range(5, ws_p.max_row):
            name = ws_p.cell(r, 2).value
            if not name or name == '合计':
                continue
                
            p_idx = r - 5
            r_att = 5 + 2 * p_idx
            r_res = 6 + 2 * p_idx
            
            a_cnt = sum(1 for c in range(4, ws_a.max_column-1) if isinstance(ws_a.cell(r_att, c).value, (int, float)))
            h_sum = sum(ws_a.cell(r_res, c).value for c in range(4, ws_a.max_column-1) if isinstance(ws_a.cell(r_res, c).value, (int, float)))
            rate = (h_sum / 8.0) / a_cnt if a_cnt > 0 else 0.0
            
            e = ws_p.cell(r, 5).value or 0.0
            f = ws_p.cell(r, 6).value or 0.0
            l = ws_p.cell(r, 12).value or 0.0
            m = ws_p.cell(r, 13).value or 0.0
            base = e + f + l + m
            alloc = base * rate
            alloc_sum += alloc
            
    diff = alloc_sum - target_fee_sum
    status = "OK (吻合)" if abs(diff) < 1.0 else f"DIFF={diff:.2f}"
    
    grand_target_all += target_fee_sum
    grand_alloc_all += alloc_sum
    
    print(f"{target_res.name[:38]:38s} | {target_fee_sum:14.2f} | {alloc_sum:14.2f} | {diff:10.2f} | {status}")

print("="*95)
print(f"{'全周期 13 季度累计总额':38s} | {grand_target_all:14.2f} | {grand_alloc_all:14.2f} | {grand_alloc_all - grand_target_all:10.2f} | 全程穿透")
print("="*95)
