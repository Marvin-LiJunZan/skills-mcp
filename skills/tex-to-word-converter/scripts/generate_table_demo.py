import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

out_dir = r"C:\JunzanLi_project\skills_mcp\skills\tex-to-word-converter\demo_output"
doc = Document()

# Set standard 1-inch margins
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# ==============================================================================
# Define dedicated Word Styles for Table Header & Table Body
# ==============================================================================
styles = doc.styles

# 1. Table Header Style: Table Header (not bold, 9.5 pt Times New Roman, centered, compact line spacing)
style_tbl_header = styles.add_style('Table Header', WD_STYLE_TYPE.PARAGRAPH)
style_tbl_header.base_style = styles['Normal']
style_tbl_header.font.name = 'Times New Roman'
style_tbl_header.font.size = Pt(9.5)
style_tbl_header.font.bold = False  # Explicitly not bold as requested
style_tbl_header.font.color.rgb = RGBColor(0, 0, 0)
style_tbl_header.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
style_tbl_header.paragraph_format.space_before = Pt(4)
style_tbl_header.paragraph_format.space_after = Pt(4)
style_tbl_header.paragraph_format.line_spacing = 1.05

# 2. Table Body Style: Table Body (not bold, 9.0 pt Times New Roman, compact line spacing)
style_tbl_body = styles.add_style('Table Body', WD_STYLE_TYPE.PARAGRAPH)
style_tbl_body.base_style = styles['Normal']
style_tbl_body.font.name = 'Times New Roman'
style_tbl_body.font.size = Pt(9.0)
style_tbl_body.font.bold = False
style_tbl_body.font.color.rgb = RGBColor(0, 0, 0)
style_tbl_body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
style_tbl_body.paragraph_format.space_before = Pt(3)
style_tbl_body.paragraph_format.space_after = Pt(3)
style_tbl_body.paragraph_format.line_spacing = 1.0

# 3. Table Footnote Style
style_tbl_note = styles.add_style('Table Note', WD_STYLE_TYPE.PARAGRAPH)
style_tbl_note.base_style = styles['Normal']
style_tbl_note.font.name = 'Times New Roman'
style_tbl_note.font.size = Pt(8.5)
style_tbl_note.font.color.rgb = RGBColor(50, 50, 50)
style_tbl_note.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
style_tbl_note.paragraph_format.space_before = Pt(3)
style_tbl_note.paragraph_format.space_after = Pt(12)

# Title & Headings
title_p = doc.add_paragraph()
r_title = title_p.add_run("Demonstration of Publication-Grade Three-Line Table in Word")
r_title.font.name = "Times New Roman"
r_title.font.size = Pt(16)
r_title.font.bold = True
title_p.paragraph_format.space_after = Pt(12)

sec_p = doc.add_paragraph()
r_sec = sec_p.add_run("2. Physical and Mechanical Properties of Specimens")
r_sec.font.name = "Times New Roman"
r_sec.font.size = Pt(13)
r_sec.font.bold = True
sec_p.paragraph_format.space_before = Pt(8)
sec_p.paragraph_format.space_after = Pt(6)

# Body paragraph with dynamic REF cross-reference
body_p = doc.add_paragraph()
r_b1 = body_p.add_run("The geotechnical and mechanical parameters obtained from comprehensive triaxial and resonant column tests are summarized in ")
r_b1.font.name = "Times New Roman"
r_b1.font.size = Pt(10.5)

# Insert REF field: Table 1
def add_ref_field(paragraph, bookmark_name, display_text):
    fldSimple = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="REF {bookmark_name} \\h"/>')
    r = parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="21"/></w:rPr><w:t>{display_text}</w:t></w:r>')
    fldSimple.append(r)
    paragraph._p.append(fldSimple)

add_ref_field(body_p, "_RefTable1", "Table 1")

r_b2 = body_p.add_run(". All specimens were consolidated under isotropic conditions prior to shearing, and the critical friction angles were calculated according to the Mohr-Coulomb criterion.")
r_b2.font.name = "Times New Roman"
r_b2.font.size = Pt(10.5)
body_p.paragraph_format.line_spacing = 1.25
body_p.paragraph_format.space_after = Pt(12)
body_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# Table Caption (Table 1: ...) with bookmark and SEQ field
cap_p = doc.add_paragraph()
cap_p.paragraph_format.space_before = Pt(8)
cap_p.paragraph_format.space_after = Pt(4)
cap_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# Bookmark start
bm_start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="101" w:name="_RefTable1"/>')
cap_p._p.append(bm_start)

# Bold "Table "
r_cap_lbl = cap_p.add_run("Table ")
r_cap_lbl.font.name = "Times New Roman"
r_cap_lbl.font.size = Pt(9.5)
r_cap_lbl.font.bold = True

# Dynamic SEQ Table field
fld_seq = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="SEQ Table \\* ARABIC"/>')
r_seq = parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:b/><w:sz w:val="19"/></w:rPr><w:t>1</w:t></w:r>')
fld_seq.append(r_seq)
cap_p._p.append(fld_seq)

# Bookmark end
bm_end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="101"/>')
cap_p._p.append(bm_end)

r_cap_dot = cap_p.add_run(". ")
r_cap_dot.font.name = "Times New Roman"
r_cap_dot.font.size = Pt(9.5)
r_cap_dot.font.bold = True

r_cap_text = cap_p.add_run("Calibrated constitutive parameters and peak mechanical properties of engineered soil-rock mixtures.")
r_cap_text.font.name = "Times New Roman"
r_cap_text.font.size = Pt(9.5)

# -------------------------------------------------------------
# Constructing Academic Standard Three-Line Table (booktabs)
# -------------------------------------------------------------
headers = [
    "Specimen ID",
    "Rock Content\n(%)",
    "Water Content\n(%)",
    "Dry Density\n(g/cm³)",
    "Cohesion c\n(kPa)",
    "Friction Angle φ\n(°)",
    "Elastic Modulus E\n(MPa)"
]

data = [
    ["SRM-00", "0.0", "12.4 ± 0.3", "1.78", "28.5 ± 1.2", "24.2 ± 0.5", "45.8 ± 2.1"],
    ["SRM-20", "20.0", "11.8 ± 0.2", "1.85", "34.2 ± 1.5", "28.7 ± 0.6", "68.4 ± 3.4"],
    ["SRM-40", "40.0", "10.5 ± 0.4", "1.96", "42.8 ± 1.8", "33.5 ± 0.8", "112.6 ± 5.2"],
    ["SRM-60", "60.0", "9.2 ± 0.3", "2.08", "51.6 ± 2.1", "38.9 ± 1.1", "185.3 ± 7.8"],
    ["SRM-80", "80.0", "8.1 ± 0.2", "2.19", "46.3 ± 2.4", "43.2 ± 1.3", "264.0 ± 9.5"],
]

table = doc.add_table(rows=len(data) + 1, cols=len(headers))
table.alignment = WD_TABLE_ALIGNMENT.CENTER

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key, val in edge_data.items():
                element.set(qn(f'w:{key}'), str(val))

# Repeat header on subsequent pages
header_tr = table.rows[0]._tr.get_or_add_trPr()
header_tr.append(OxmlElement('w:tblHeader'))

# Fill Header Row (Using 'Table Header' style, bold = False)
for c_idx, text in enumerate(headers):
    cell = table.rows[0].cells[c_idx]
    p = cell.paragraphs[0]
    p.style = doc.styles['Table Header']
    run = p.add_run(text)
    
    # 1.5 pt top border, 0.75 pt mid border, 0 vertical borders
    set_cell_border(cell, 
                    top={'val': 'single', 'sz': '12', 'color': '000000'},
                    bottom={'val': 'single', 'sz': '6', 'color': '000000'},
                    left={'val': 'none'}, right={'val': 'none'},
                    insideH={'val': 'none'}, insideV={'val': 'none'})

# Fill Data Rows (Using 'Table Body' style)
num_rows = len(table.rows)
for r_idx, row_data in enumerate(data, start=1):
    row = table.rows[r_idx]
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    
    is_last = (r_idx == num_rows - 1)
    for c_idx, val in enumerate(row_data):
        cell = row.cells[c_idx]
        p = cell.paragraphs[0]
        p.style = doc.styles['Table Body']
        run = p.add_run(val)
        
        # Last row 1.5 pt bottom border, internal rows none
        if is_last:
            set_cell_border(cell,
                            top={'val': 'none'},
                            bottom={'val': 'single', 'sz': '12', 'color': '000000'},
                            left={'val': 'none'}, right={'val': 'none'},
                            insideH={'val': 'none'}, insideV={'val': 'none'})
        else:
            set_cell_border(cell,
                            top={'val': 'none'}, bottom={'val': 'none'},
                            left={'val': 'none'}, right={'val': 'none'},
                            insideH={'val': 'none'}, insideV={'val': 'none'})

# Footnote with 'Table Note' style
note_p = doc.add_paragraph()
note_p.style = doc.styles['Table Note']
r_note_lbl = note_p.add_run("Note: ")
r_note_lbl.font.italic = True
r_note_txt = note_p.add_run("Values represent mean ± standard deviation from five repeated triaxial shearing runs under consolidated drained (CD) conditions; c and φ denote effective shear parameters.")

out_file = os.path.join(out_dir, "Publication_Grade_Three_Line_Table_Demo.docx")
doc.save(out_file)
print(f"[SUCCESS] Updated three-line table demo saved at: {out_file}")
