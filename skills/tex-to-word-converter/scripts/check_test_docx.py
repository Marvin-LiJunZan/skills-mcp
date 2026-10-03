import docx
from docx import Document

doc = Document(r"C:\JunzanLi_project\skills_mcp\skills\tex-to-word-converter\demo_output\test_math.docx")
p0 = doc.paragraphs[0]
print("Paragraph 0 text:", p0.text)
print("Paragraph 0 runs count:", len(p0.runs))
for i, r in enumerate(p0.runs):
    font_style = []
    if r.font.italic: font_style.append("italic")
    if r.font.subscript: font_style.append("subscript")
    if r.font.superscript: font_style.append("superscript")
    style_desc = f" [{', '.join(font_style)}]" if font_style else ""
    print(f"  Run {i}: '{r.text}'{style_desc}")

print("\nParagraph 0 XML snippet:")
print(p0._p.xml[:500])

print("\nParagraph 1 (Display Equation):")
print(doc.paragraphs[1]._p.xml[:400])
