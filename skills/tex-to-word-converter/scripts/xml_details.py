import docx
from docx import Document
import re

doc_path = r"C:\JunzanLi_project\skills_mcp\skills\tex-to-word-converter\demo_output\Publication_Grade_2x2_Subfigures_Demo.docx"
doc = Document(doc_path)

p2_xml = doc.paragraphs[2]._p.xml
instrs = re.findall(r'<w:instrText[^>]*>(.*?)</w:instrText>', p2_xml)
print("P2 instrText:", instrs)

t_xml = doc.tables[0]._tbl.xml
t_instrs = re.findall(r'<w:instrText[^>]*>(.*?)</w:instrText>', t_xml)
print("Table instrText:", t_instrs)

# Check gridSpan or merge in row 4
r4_xml = doc.tables[0].rows[4]._tr.xml
print("Row 4 gridSpan/merge:", re.findall(r'<w:gridSpan[^>]*/>', r4_xml), re.findall(r'<w:hMerge[^>]*/>', r4_xml))
