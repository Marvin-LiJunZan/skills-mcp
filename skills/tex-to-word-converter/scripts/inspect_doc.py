import docx
from docx import Document

doc_path = r"C:\JunzanLi_project\skills_mcp\skills\tex-to-word-converter\demo_output\Publication_Grade_2x2_Subfigures_Demo.docx"
doc = Document(doc_path)

print("=== PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"P[{i}]: {p.text}")

print("\n=== TABLES ===")
for t_idx, table in enumerate(doc.tables):
    print(f"Table {t_idx} (rows={len(table.rows)}, cols={len(table.columns)}):")
    for r_idx, row in enumerate(table.rows):
        row_texts = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        print(f"  Row {r_idx}: {row_texts}")
