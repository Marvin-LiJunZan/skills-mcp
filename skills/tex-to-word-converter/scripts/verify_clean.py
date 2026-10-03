from docx import Document
from docx.oxml.ns import qn

doc = Document(r"C:\JunzanLi_project\skills_mcp\skills\tex-to-word-converter\demo_output\test_math_clean.docx")
p0 = doc.paragraphs[0]
direct_omaths = p0._p.findall(qn('m:oMath'))
print("P0 direct inline m:oMath count:", len(direct_omaths))

p1 = doc.paragraphs[1]
p1_omathparas = p1._p.findall(qn('m:oMathPara'))
print("P1 display m:oMathPara count (kept for MathType):", len(p1_omathparas))
