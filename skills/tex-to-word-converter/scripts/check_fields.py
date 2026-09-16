import docx
from docx import Document

doc_path = r"C:\JunzanLi_project\skills_mcp\skills\tex-to-word-converter\demo_output\Publication_Grade_2x2_Subfigures_Demo.docx"
doc = Document(doc_path)

# Check paragraph 2 XML for field codes
p2 = doc.paragraphs[2]
print("P2 XML snippet:")
print(p2._p.xml[:1000])

# Check Table Row 4 XML
t = doc.tables[0]
r4 = t.rows[4]
print("\nRow 4 Cell 0 XML snippet:")
print(r4.cells[0]._tc.xml[:1000])
