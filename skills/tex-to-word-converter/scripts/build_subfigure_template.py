import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

out_path = r'C:\JunzanLi_project\skills_mcp\skills\tex-to-word-converter\templates\subfigure_template.docx'
doc = Document()

# Page Margins: Standard Elsevier/Academic 1 inch
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

def remove_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(r'''
        <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="none"/>
            <w:left w:val="none"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="none"/>
            <w:insideV w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

def add_field(paragraph, instr_text, display_text, font_size=19, bold=True):
    fld = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="{instr_text}"/>')
    b_xml = '<w:b/>' if bold else ''
    r = parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>{b_xml}<w:sz w:val="{font_size}"/></w:rPr><w:t>{display_text}</w:t></w:r>')
    fld.append(r)
    paragraph._p.append(fld)

demo_dir = r'C:\JunzanLi_project\skills_mcp\skills\tex-to-word-converter\demo_output'
fig_a = os.path.join(demo_dir, 'sub_a.png')
fig_b = os.path.join(demo_dir, 'sub_b.png')
fig_c = os.path.join(demo_dir, 'sub_c.png')
fig_d = os.path.join(demo_dir, 'sub_d.png')

# Header explanation
p0 = doc.add_paragraph('Standard Subfigure Templates (Reference Ground Truth for AI & Converters)')
p0.runs[0].font.bold = True
p0.runs[0].font.size = Pt(14)

# ----------------- 1. Standard 2x2 Subfigure Panel -----------------
p_h1 = doc.add_paragraph('1. Canonical 2x2 Subfigure Panel (4 Panels, GridSpan=2)')
p_h1.runs[0].font.bold = True
p_h1.runs[0].font.size = Pt(11)

t2x2 = doc.add_table(rows=5, cols=2)
t2x2.alignment = WD_TABLE_ALIGNMENT.CENTER
for r in t2x2.rows:
    for c in r.cells:
        c.width = Inches(3.15)
        remove_borders(c)

# Row 0: Images A, B
if os.path.exists(fig_a):
    p = t2x2.rows[0].cells[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(fig_a, width=Inches(3.05))
if os.path.exists(fig_b):
    p = t2x2.rows[0].cells[1].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(fig_b, width=Inches(3.05))

# Row 1: Labels A, B
p = t2x2.rows[1].cells[0].paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(8)
r = p.add_run('(a) Deviatoric stress versus axial strain curves')
r.font.name = 'Times New Roman'
r.font.size = Pt(9.5)

p = t2x2.rows[1].cells[1].paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(8)
r = p.add_run('(b) Excess pore pressure dissipation response')
r.font.name = 'Times New Roman'
r.font.size = Pt(9.5)

# Row 2: Images C, D
if os.path.exists(fig_c):
    p = t2x2.rows[2].cells[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(fig_c, width=Inches(3.05))
if os.path.exists(fig_d):
    p = t2x2.rows[2].cells[1].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(fig_d, width=Inches(3.05))

# Row 3: Labels C, D
p = t2x2.rows[3].cells[0].paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(6)
r = p.add_run('(c) Kinematic yield locus expansion in p′-q space')
r.font.name = 'Times New Roman'
r.font.size = Pt(9.5)

p = t2x2.rows[3].cells[1].paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(6)
r = p.add_run('(d) Real-time sensor monitoring displacement tracking')
r.font.name = 'Times New Roman'
r.font.size = Pt(9.5)

# Row 4: Fused bottom caption
cell_cap = t2x2.rows[4].cells[0]
tcPr_cap = cell_cap._tc.get_or_add_tcPr()
tcPr_cap.append(parse_xml(f'<w:gridSpan {nsdecls("w")} w:val="2"/>'))
p_cap = cell_cap.paragraphs[0]
p_cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p_cap.paragraph_format.space_before = Pt(6)
p_cap.paragraph_format.space_after = Pt(12)

# Bookmark strictly wraps 'Fig. 1'
p_cap._p.append(parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="501" w:name="_RefFig1"/>'))
p_cap.add_run('Fig. ').bold = True
add_field(p_cap, 'SEQ Fig. \\* ARABIC', '1', font_size=19, bold=True)
p_cap._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="501"/>'))

r_cap_txt = p_cap.add_run('. Multi-dimensional constitutive and transport validation: (a) Stress-strain; (b) Pore water pressure; (c) Yield locus; (d) Sensor tracking.')
r_cap_txt.font.name = 'Times New Roman'
r_cap_txt.font.size = Pt(9.5)

# ----------------- 2. Standard 1x2 Subfigure Panel -----------------
p_h2 = doc.add_paragraph('2. Canonical 1x2 Subfigure Panel (2 Panels Side-by-Side, GridSpan=2)')
p_h2.runs[0].font.bold = True
p_h2.runs[0].font.size = Pt(11)

t1x2 = doc.add_table(rows=3, cols=2)
t1x2.alignment = WD_TABLE_ALIGNMENT.CENTER
for r in t1x2.rows:
    for c in r.cells:
        c.width = Inches(3.15)
        remove_borders(c)

# Row 0: Images A, B
if os.path.exists(fig_a):
    p = t1x2.rows[0].cells[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(fig_a, width=Inches(3.05))
if os.path.exists(fig_b):
    p = t1x2.rows[0].cells[1].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(fig_b, width=Inches(3.05))

# Row 1: Labels A, B
p = t1x2.rows[1].cells[0].paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(8)
r = p.add_run('(a) Measured versus computed deviatoric stress response')
r.font.name = 'Times New Roman'
r.font.size = Pt(9.5)

p = t1x2.rows[1].cells[1].paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(8)
r = p.add_run('(b) Normalized excess pore pressure dissipation trends')
r.font.name = 'Times New Roman'
r.font.size = Pt(9.5)

# Row 2: Fused bottom caption
cell_cap2 = t1x2.rows[2].cells[0]
tcPr_cap2 = cell_cap2._tc.get_or_add_tcPr()
tcPr_cap2.append(parse_xml(f'<w:gridSpan {nsdecls("w")} w:val="2"/>'))
p_cap2 = cell_cap2.paragraphs[0]
p_cap2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p_cap2.paragraph_format.space_before = Pt(6)
p_cap2.paragraph_format.space_after = Pt(12)

# Bookmark strictly wraps 'Fig. 2'
p_cap2._p.append(parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="502" w:name="_RefFig2"/>'))
p_cap2.add_run('Fig. ').bold = True
add_field(p_cap2, 'SEQ Fig. \\* ARABIC', '2', font_size=19, bold=True)
p_cap2._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="502"/>'))

r_cap_txt2 = p_cap2.add_run('. Side-by-side comparison of constitutive behaviors: (a) Stress-strain; (b) Dissipation rates.')
r_cap_txt2.font.name = 'Times New Roman'
r_cap_txt2.font.size = Pt(9.5)

doc.save(out_path)
print(f'[SUCCESS] Generated subfigure template: {out_path}')
