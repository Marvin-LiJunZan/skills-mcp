# -*- coding: utf-8 -*-
"""
工程科技申报材料一致性与去AI味自动化核验脚本
Audit Consistency & De-AI Verification for Engineering Grant Proposals
"""

import os
import re
import sys
import docx

# 高频AI套话黑名单
AI_BUZZWORDS = [
    "具有十分重要的意义", "具有极其重要的意义", "具有深远的战略意义",
    "显著提升了", "取得了突破性进展", "取得了里程碑式",
    "奠定了坚实的基础", "开创了全新局面", "不可否认的是",
    "总而言之", "综上所述，我们可以看出", "毫无疑问"
]

# 防御性/心虚措辞黑名单
DEFENSIVE_PATTERNS = [
    "由于试验条件有限", "试验条件所限", "本研究尚未考虑",
    "仅做初步探索", "存在一定的不确定性", "可能存在局限性",
    "尚待进一步验证", "难免存在不足"
]

# 专利号正则表达式
PATENT_REGEX = re.compile(r'ZL\s*20\d{2}[128]\d{7}\.[\dX]', re.IGNORECASE)

def audit_docx(filepath):
    print(f"\n==========================================")
    print(f"开始审查文档: {os.path.basename(filepath)}")
    print(f"==========================================")
    
    if not os.path.exists(filepath):
        print(f"错误: 文件不存在 -> {filepath}")
        return

    doc = docx.Document(filepath)
    full_text = []
    
    # 提取段落
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            full_text.append(t)
            
    # 提取表格文本
    table_cell_texts = []
    for t_idx, t in enumerate(doc.tables):
        for r_idx, r in enumerate(t.rows):
            for c_idx, c in enumerate(r.cells):
                txt = c.text.strip()
                if txt:
                    table_cell_texts.append((t_idx+1, r_idx+1, c_idx+1, txt))
                    full_text.append(txt)

    combined_text = "\n".join(full_text)
    
    # 1. 检查 AI 套话
    ai_found = []
    for bw in AI_BUZZWORDS:
        count = combined_text.count(bw)
        if count > 0:
            ai_found.append((bw, count))
            
    if ai_found:
        print("[警告] 检出 AI 高频空泛套话:")
        for bw, count in ai_found:
            print(f"  - 词汇: '{bw}' (出现 {count} 次) -> 建议替换为实测数据或工艺参数")
    else:
        print("[通过] 未检出常见 AI 高频空泛套话。")

    # 2. 检查防御性心虚措辞
    def_found = []
    for dp in DEFENSIVE_PATTERNS:
        count = combined_text.count(dp)
        if count > 0:
            def_found.append((dp, count))
            
    if def_found:
        print("\n[警告] 检出防御性/心虚语气词 (违背反防御写作规范):")
        for dp, count in def_found:
            print(f"  - 词汇: '{dp}' (出现 {count} 次) -> 建议删除或改写为客观适用范围描述")
    else:
        print("[通过] 未检出防御性心虚措辞。")

    # 3. 提取并核对专利号
    patents = PATENT_REGEX.findall(combined_text)
    unique_patents = sorted(list(set(patents)))
    print(f"\n[信息] 文档中识别到的专利号 ({len(unique_patents)} 个):")
    for pat in unique_patents:
        print(f"  - {pat}")

    # 4. 关键易错数值警示
    if "26人" in combined_text:
        print("\n[严重错误警示] 文档中包含 '26人'！请复核：历史决议表中26为专家评审得分，切勿误记为研发人员人数！")
    else:
        print("[通过] 未发现 '26人' 历史混淆错误。")

    print(f"\n文档审查完毕: {os.path.basename(filepath)}")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        for arg in sys.argv[1:]:
            audit_docx(arg)
    else:
        print("用法: python audit_grant_consistency.py <文件路径.docx> ...")
