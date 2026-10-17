# -*- coding: utf-8 -*-
"""
多源财务数据穿透提取模块 (Multi-source Financial Data Loader)
通用自动提取：
1. 任务书 docx 中的课题人员名单与层级角色
2. 各月工资表（应发合计）
3. 各月社保申报表（6险/5险单位缴存额）
4. 各月住房公积金申报表（单位缴存额，缺失时回退读取工资表公积金列）
5. 各月研发费台账（动态匹配课题列与合计行，提取法定目标总额）
"""
import os, glob, re
import docx, xlrd, openpyxl

def parse_task_book(docx_path, target_unit="总承包"):
    """
    解析科研课题合同任务书 docx，提取牵头单位/指定单位的研究人员名单及职务
    """
    if not os.path.exists(docx_path):
        return []
    doc = docx.Document(docx_path)
    # 查找人员名单表格（通常包含：负责人、职务职称、主要研究人员等）
    target_table = None
    for t in doc.tables:
        txt = "".join([c.text.replace("\n", "").replace(" ", "") for row in t.rows[:5] for c in row.cells])
        if "主要研究人员" in txt or ("负责人" in txt and "职务" in txt):
            target_table = t
            break
            
    if not target_table:
        return []

    people = []
    for r_idx in range(5, len(target_table.rows)):
        row = target_table.rows[r_idx]
        cells = [c.text.strip().replace("\n", "").replace(" ", "") for c in row.cells]
        if len(cells) >= 6:
            name = cells[0]
            role = cells[3]
            unit = cells[5]
            if name and name not in ['负责人', '主要研究人员', '姓名', '序号', '']:
                if not target_unit or target_unit in unit:
                    people.append({
                        'name': name,
                        'role': role,
                        'unit': unit,
                        'tier': classify_role(role, name)
                    })
    return people

def classify_role(role_str, name=""):
    """
    根据职务名称与人员姓名智能归类为三档工时率级别：
    - leader: 领导层，建议工时率 0.25
    - key: 核心骨干层，建议工时率 0.15 ~ 0.20
    - general: 普通研发人员，建议工时率 0.10
    """
    if name == "陈前": # 课题负责人
        return 'leader'
    r = role_str.replace(" ", "")
    if any(k in r for k in ['总经理', '总工程师', '副总经理', '总工', '副总工', '主任', '负责人', '安全总监']):
        return 'leader'
    elif any(k in r for k in ['副组长', '部长', '副部长', '骨干', '组长', '项目经理', '高级工程师']):
        return 'key'
    else:
        return 'general'

def load_monthly_wage(payroll_dir, year, month):
    """
    从 {year}年{month}月份工资表（总包部）.xlsx 提取应发工资与公积金备份
    """
    files = [
        os.path.join(payroll_dir, f) for f in os.listdir(payroll_dir)
        if '工资' in f and not f.startswith("~$") and re.search(rf"{year}年{month}月份?(?!\d)", f)
    ]
    if not files:
        return {}
    
    target_file = files[0]
    wb = openpyxl.load_workbook(target_file, data_only=True)
    ws = wb['汇总工资表'] if '汇总工资表' in wb.sheetnames else wb.active
    
    # 动态定位姓名列、应发合计列、公积金列
    name_col = 2
    wage_col = 24
    gjj_col = 28
    for r in range(1, min(6, ws.max_row+1)):
        for c in range(1, min(35, ws.max_column+1)):
            v = str(ws.cell(r, c).value or "").replace("\n", "").strip()
            if v == "姓名": name_col = c
            elif "应发合计" in v or "应发工资" in v: wage_col = c
            elif "住房公积金" in v or "公积金" in v: gjj_col = c
            
    res = {}
    for r in range(4, ws.max_row+1):
        name = str(ws.cell(r, name_col).value or "").strip()
        wage = ws.cell(r, wage_col).value
        gjj = ws.cell(r, gjj_col).value
        if name and name not in ['姓名', '合计', '小计', 'None']:
            try:
                w_val = float(wage) if wage not in [None, ''] else 0.0
                g_val = float(gjj) if gjj not in [None, ''] else 0.0
                res[name] = {'应发工资': round(w_val, 2), '公积金备份': round(g_val, 2)}
            except (ValueError, TypeError):
                pass
    return res

def load_monthly_social_security(insurance_dir, year, month):
    """
    从 社会保险费申报个人明细表-{year}年{month}月.xls 提取单位各险种
    """
    year_dir = os.path.join(insurance_dir, f"{year}年") if os.path.exists(os.path.join(insurance_dir, f"{year}年")) else insurance_dir
    candidates = [
        os.path.join(year_dir, f) for f in os.listdir(year_dir)
        if '社会保险' in f and not f.startswith("~$") and "调整" not in f and "新增" not in f and re.search(rf"{year}年{month}月(?!\d)", f)
    ]
    if not candidates:
        return {}, False

    target_file = candidates[0]
    is_xlsx = target_file.endswith('.xlsx')
    
    data = {}
    has_maternity = False
    
    if is_xlsx:
        wb = openpyxl.load_workbook(target_file, data_only=True)
        sname = None
        for s in wb.sheetnames:
            if any(k in s for k in ['人员', '明细', f'{month}月']):
                sname = s
                break
        if not sname:
            sname = max(wb.sheetnames, key=lambda s: wb[s].max_row)
        ws = wb[sname]
        # 解析表头
        hdr6 = [str(ws.cell(6, c).value or "") for c in range(1, ws.max_column+1)]
        has_maternity = any('生育' in x for x in hdr6)
        # 读取行
        for r in range(8, ws.max_row+1):
            name = str(ws.cell(r, 1).value or "").strip()
            if name and name not in ['姓名', '合计', 'None']:
                # F:养老(6), I:工伤(9), L:失业(12), O:医疗(15), R:大病(18), U:生育(21), W:单位合计(23)
                def gv(col_idx):
                    val = ws.cell(r, col_idx).value
                    try: return round(float(val), 2) if val not in [None, ''] else 0.0
                    except: return 0.0
                data[name] = {
                    '养老': gv(6),
                    '工伤': gv(9),
                    '失业': gv(12),
                    '医疗': gv(15),
                    '大病': gv(18),
                    '生育': gv(21) if has_maternity else 0.0,
                    '单位合计': gv(23)
                }
    else:
        wb = xlrd.open_workbook(target_file)
        sname = None
        for s in wb.sheet_names():
            if any(k in s for k in ['人员', '明细', f'{month}月']):
                sname = s
                break
        if not sname:
            sname = max(wb.sheet_names(), key=lambda s: wb.sheet_by_name(s).nrows)
        ws = wb.sheet_by_name(sname)
        all_hdr = "".join([str(ws.cell_value(r, c)) for r in range(min(8, ws.nrows)) for c in range(ws.ncols)])
        has_maternity = '生育' in all_hdr
        for r in range(7, ws.nrows):
            name = str(ws.cell_value(r, 0) or "").strip()
            if name and name not in ['姓名', '合计', 'None', '']:
                def gv(col_idx):
                    val = ws.cell_value(r, col_idx) if col_idx < ws.ncols else 0.0
                    try: return round(float(val), 2) if val not in [None, ''] else 0.0
                    except: return 0.0
                data[name] = {
                    '养老': gv(5),
                    '工伤': gv(8),
                    '失业': gv(11),
                    '医疗': gv(14),
                    '大病': gv(17),
                    '生育': gv(20) if has_maternity else 0.0,
                    '单位合计': gv(22)
                }
                
    return data, has_maternity

def load_monthly_housing_fund(insurance_dir, year, month, backup_wages=None):
    """
    从 住房公积金申报个人明细表-{year}年{month}月.xlsx 提取单位缴存额
    如果该月公积金申报表缺失，则自动使用工资表备份中的公积金列
    """
    year_dir = os.path.join(insurance_dir, f"{year}年") if os.path.exists(os.path.join(insurance_dir, f"{year}年")) else insurance_dir
    candidates = [
        os.path.join(year_dir, f) for f in os.listdir(year_dir)
        if '公积金' in f and not f.startswith("~$") and "调整" not in f and re.search(rf"{year}年{month}月(?!\d)", f)
    ]
    
    res = {}
    if candidates:
        target_file = candidates[0]
        wb = openpyxl.load_workbook(target_file, data_only=True)
        sname = None
        for s in wb.sheetnames:
            if any(k in s for k in ['人员', '明细', f'{month}月']):
                sname = s
                break
        if not sname:
            sname = max(wb.sheetnames, key=lambda s: wb[s].max_row)
        ws = wb[sname]
        name_col = 3
        gjj_col = 7
        for r in range(1, min(5, ws.max_row+1)):
            for c in range(1, min(15, ws.max_column+1)):
                v = str(ws.cell(r, c).value or "").replace("\n", "").strip()
                if v == "姓名": name_col = c
                elif "单位缴存额" in v: gjj_col = c
                
        for r in range(3, ws.max_row+1):
            name = str(ws.cell(r, name_col).value or "").strip()
            val = ws.cell(r, gjj_col).value
            if name and name not in ['姓名', '合计', '小计', 'None']:
                try:
                    res[name] = round(float(val), 2) if val not in [None, ''] else 0.0
                except:
                    pass
    elif backup_wages:
        # 回退使用工资表备份
        for name, w_info in backup_wages.items():
            if '公积金备份' in w_info and w_info['公积金备份'] > 0:
                res[name] = w_info['公积金备份']
                
    return res

def load_monthly_ledger_target(ledger_dir, year, month, project_keywords):
    """
    从 研发费职工薪酬分配表 -{year}.xls 动态扫描课题关键字与合计行，提取该月的研发工资目标金额
    """
    fpath = os.path.join(ledger_dir, f"研发费职工薪酬分配表 -{year}.xls")
    if not os.path.exists(fpath):
        return 0.0, None, None
        
    wb = xlrd.open_workbook(fpath)
    sname = f"{year}年{month}月工资"
    if sname not in wb.sheet_names():
        return 0.0, None, None
        
    ws = wb.sheet_by_name(sname)
    proj_col = None
    for c in range(ws.ncols):
        txt = " ".join([str(ws.cell_value(r, c)).replace("\n", " ") for r in range(min(5, ws.nrows))])
        if any(k in txt for k in project_keywords):
            proj_col = c
            break
            
    if proj_col is None:
        return 0.0, None, None
        
    fee_col = proj_col + 1 # 研发工资列
    tot_row = None
    for r in range(ws.nrows):
        if str(ws.cell_value(r, 1)).strip() == '合计':
            tot_row = r
            break
            
    if tot_row is not None:
        target_val = ws.cell_value(tot_row, fee_col)
        try:
            return round(float(target_val), 2), proj_col, tot_row
        except:
            return 0.0, proj_col, tot_row
            
    return 0.0, proj_col, None

if __name__ == '__main__':
    # 模块自测：测试 2018 年 1~3 月多源数据完整度
    tb_path = r"C:\JunzanLi_project\课题申报横向\财务\19-B-03科研课题合同任务书_1005凝练版.docx"
    payroll_dir = r"C:\JunzanLi_project\课题申报横向\财务\18和22号线2018-2021"
    insurance_dir = r"C:\JunzanLi_project\课题申报横向\财务\2018-2021社保公积金等"
    ledger_dir = r"C:\JunzanLi_project\课题申报横向\财务\1822号线研发薪酬分摊台账"
    
    people = parse_task_book(tb_path)
    print(f"成功提取任务书总承包部人员: {len(people)} 人")
    for p in people[:5]:
        print(" ", p)
        
    for m in [1, 2, 3]:
        w = load_monthly_wage(payroll_dir, 2018, m)
        sb, has_mat = load_monthly_social_security(insurance_dir, 2018, m)
        gjj = load_monthly_housing_fund(insurance_dir, 2018, m, w)
        tgt, pcol, trow = load_monthly_ledger_target(ledger_dir, 2018, m, ["智能设计"])
        print(f"\n2018年{m}月: 工资覆盖={len(w)}人, 社保覆盖={len(sb)}人(含生育={has_mat}), 公积金覆盖={len(gjj)}人, 台账目标={tgt:.2f} (行{trow+1})")
        # 打印陈前数据验证
        if "陈前" in w:
            cq_w = w["陈前"]['应发工资']
            cq_sb = sb.get("陈前", {})
            cq_gjj = gjj.get("陈前", 0.0)
            print(f"  [陈前核验] 应发={cq_w:.2f}, 社保合计={cq_sb.get('单位合计',0):.2f} (养老={cq_sb.get('养老',0)}, 生育={cq_sb.get('生育',0)}), 公积金={cq_gjj:.2f}")
