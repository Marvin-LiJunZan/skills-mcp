# -*- coding: utf-8 -*-
"""
19-B-03 专属生成器（2019年3月首月样板验证）：
- 项目名称：地铁工程总承包模式下如何应用BIM平台提高项目管控效率的研究
- 编号：19-B-03
- 负责人：陈前
- 月度目标：最后一列小计 59,880.74 元（约6万元）
- 工时率：领导 0.25，骨干 0.15~0.20，普通 0.10
- 出勤打点：8, 4, 2, 1, 0.5, /
- 严格基于 21-B-02 母版克隆与样式复刻
"""
import os, sys, datetime as dt, re
import openpyxl, xlrd
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

sys.path.append(r"C:\JunzanLi_project\skills_mcp\skills\rd-multisource-payroll-allocator\scripts")
from data_loader import (
    parse_task_book,
    load_monthly_wage,
    load_monthly_social_security,
    load_monthly_housing_fund
)

tpl_path = r"C:\JunzanLi_project\课题申报横向\人工费考勤表与分摊表\2024年7-9月研发人员薪酬、职工薪酬分配表21-B-02（定稿打印版）.xlsx"
tb_path = r"C:\JunzanLi_project\课题申报横向\财务\19-B-03科研课题合同任务书_1005凝练版.docx"
payroll_dir = r"C:\JunzanLi_project\课题申报横向\财务\18和22号线2018-2021"
insurance_dir = r"C:\JunzanLi_project\课题申报横向\财务\2018-2021社保公积金等"

out_dir = r"C:\JunzanLi_project\课题申报横向\人工费考勤表与分摊表\结果\Antigravity_批量生成结果_19-B-03\样板测试_2019年3月"
os.makedirs(out_dir, exist_ok=True)
out_file = os.path.join(out_dir, "2019年3月研发人员薪酬、职工薪酬分配表19-B-03.xlsx")

year, month = 2019, 3
monthly_target = 59880.74 # 3,113,798.60 / 52

print(f">>> [1] 加载 {year}年{month}月 财务基础数据...")
tb_people = parse_task_book(tb_path)
tb_map = {p['name']: p for p in tb_people}
wages = load_monthly_wage(payroll_dir, year, month)
ss_dict, has_mat = load_monthly_social_security(insurance_dir, year, month)
hf_dict = load_monthly_housing_fund(insurance_dir, year, month, wages)

print(f"  任务书总包部名单: {len(tb_people)} 人")
print(f"  工资表覆盖: {len(wages)} 人, 社保: {len(ss_dict)} 人, 公积金: {len(hf_dict)} 人")

# 确定当期参研人员（任务书总包部在职人员）
active_people = [p for p in tb_people if p['name'] in wages]
print(f"  当月在职参研人员: {len(active_people)} 人")

# 计算考勤周期：2019-02-21 至 2019-03-20
start_date = dt.date(2019, 2, 21)
end_date = dt.date(2019, 3, 20)
days_count = (end_date - start_date).days + 1
days_range = [start_date + dt.timedelta(days=i) for i in range(days_count)]

def get_mark(d):
    # 周期内无元旦或春节（2019年春节为2月4-10日，已过）
    if d.weekday() >= 5: # 周末
        return "/"
    return 8

marks = [get_mark(d) for d in days_range]
workday_indices = [i for i, m in enumerate(marks) if m == 8]
W_days = len(workday_indices)
print(f"  考勤周期: {start_date} ~ {end_date} (共 {days_count} 天, 额定工作日: {W_days} 天)")

# 逆向求解工时与工时率配平（目标：最后一列小计等于 monthly_target）
# 为每位人员计算管理费用小计基数 base_fee = 工资 + 社保 + 公积金
person_base_fees = {}
assigned_plans = {}
for p in active_people:
    nm = p['name']
    w = wages[nm]['应发工资']
    sb = ss_dict.get(nm, {}).get('单位合计', 0.0)
    hf = hf_dict.get(nm, round(w * 0.12, 2))
    base_fee = w + sb + hf
    person_base_fees[nm] = base_fee
    
    tier = p['tier']
    if tier == 'leader':
        base_rate = 0.25
    elif tier == 'key':
        base_rate = 0.18
    else:
        base_rate = 0.10
    assigned_plans[nm] = {
        'name': nm,
        'wage': w,
        'base_fee': base_fee,
        'tier': tier,
        'base_rate': base_rate
    }

init_sum = sum(p['base_fee'] * p['base_rate'] for p in assigned_plans.values())
scale_factor = monthly_target / init_sum
print(f"  初始计划总额: {init_sum:.2f} 元, 缩放系数: {scale_factor:.4f}")

for nm, p in assigned_plans.items():
    p['eff_rate'] = p['base_rate'] * scale_factor
    # 转换为考勤小时 (允许 0.5 小时网格)
    p['h_hours'] = round(p['eff_rate'] * 8.0 * W_days * 2) / 2
    p['calc_rate'] = (p['h_hours'] / 8.0) / W_days
    p['calc_fee'] = p['base_fee'] * p['calc_rate']

calc_tot = sum(p['calc_fee'] for p in assigned_plans.values())
diff = monthly_target - calc_tot
print(f"  网格离散实算小计总额: {calc_tot:.2f} 元, 尾差: {diff:+.2f} 元")

first_person = active_people[0]['name'] # 陈前
assigned_plans[first_person]['comp'] = round(diff, 2)
print(f"  在课题负责人 ({first_person}) 研发工资公式注入尾差微调: {diff:+.2f} 元, 达成最后一列绝对 0.00 元平衡!")

print("\n>>> [2] 克隆母版并生成 2019年3月 工作簿...")
wb_out = openpyxl.load_workbook(tpl_path)
wb_out['202407月'].title = '201903月'
wb_out['202407月考勤表'].title = '201903月考勤表'
for s in ['202408月', '202408月考勤表', '202409月', '202409月考勤表']:
    if s in wb_out.sheetnames:
        del wb_out[s]

ws_pay = wb_out['201903月']
ws_att = wb_out['201903月考勤表']

font_body = Font(name="宋体", size=10)
font_bold = Font(name="宋体", size=10, bold=True)
align_center = Alignment(horizontal="center", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
thin_border = Side(border_style="thin", color="000000")
border_all = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

# ----------------- 改造分摊表 (ws_pay) 插入 2.6生育保险 列 -----------------
for m in list(ws_pay.merged_cells.ranges):
    if m.min_row <= 4:
        ws_pay.unmerge_cells(range_string=str(m))

ws_pay.insert_cols(12, 1) # Col L
ws_pay.insert_cols(23, 1) # Col W

ws_pay.merge_cells("B1:Z1")
ws_pay.cell(1, 2, value="研发费用—职工薪酬分摊明细表").font = Font(name="宋体", size=22, bold=True)
ws_pay.cell(1, 2).alignment = align_center

ws_pay.merge_cells("A2:M2")
ws_pay.cell(2, 1, value="项目名称：地铁工程总承包模式下如何应用BIM平台提高项目管控效率的研究").font = font_body

ws_pay.merge_cells("N2:P2")
ws_pay.cell(2, 14, value="课题合同（项目）编号：").font = font_body

ws_pay.merge_cells("Q2:V2")
ws_pay.cell(2, 17, value="19-B-03").font = font_body

ws_pay.merge_cells("W2:Z2")
ws_pay.cell(2, 23, value=f"{year}年{month}月（单位：元）").font = font_body
ws_pay.cell(2, 23).alignment = align_right

ws_pay.merge_cells("A3:A4"); ws_pay.cell(3, 1, value="序号").font = font_body; ws_pay.cell(3, 1).alignment = align_center
ws_pay.merge_cells("B3:B4"); ws_pay.cell(3, 2, value="姓名").font = font_body; ws_pay.cell(3, 2).alignment = align_center
ws_pay.merge_cells("C3:C4"); ws_pay.cell(3, 3, value="从事研发天数").font = font_body; ws_pay.cell(3, 3).alignment = align_center
ws_pay.merge_cells("D3:D4"); ws_pay.cell(3, 4, value="工时率").font = font_body; ws_pay.cell(3, 4).alignment = align_center

ws_pay.merge_cells("E3:O3"); ws_pay.cell(3, 5, value="管理费用").font = font_body; ws_pay.cell(3, 5).alignment = align_center
ws_pay.merge_cells("P3:Z3"); ws_pay.cell(3, 16, value="计入课题的研发人员工资性费用").font = font_body; ws_pay.cell(3, 16).alignment = align_center

col_titles_mgt = ["1.应发工资", "2.社会保险费", "2.1基本养老", "2.2工伤", "2.3失业", "2.4基本医疗", "2.5大病医疗", "2.6生育保险", "3.住房\n公积金", "4.企业年金", "小计"]
for i, t in enumerate(col_titles_mgt):
    c = ws_pay.cell(4, 5 + i, value=t)
    c.font = font_body; c.alignment = align_center; c.border = border_all

col_titles_rd = ["1.研发工资", "2.社会保险费", "2.1基本养老", "2.2工伤", "2.3失业", "2.4基本医疗", "2.5大病医疗", "2.6生育保险", "3.住房\n公积金", "4.企业年金", "小计"]
for i, t in enumerate(col_titles_rd):
    c = ws_pay.cell(4, 16 + i, value=t)
    c.font = font_body; c.alignment = align_center; c.border = border_all

for c in range(1, 27):
    ws_pay.cell(3, c).border = border_all

ws_pay.column_dimensions['L'].width = 11.5
ws_pay.column_dimensions['W'].width = 11.5

if ws_pay.max_column > 26:
    ws_pay.delete_cols(27, ws_pay.max_column - 26 + 5)

for m in list(ws_pay.merged_cells.ranges):
    if m.min_row >= 5:
        ws_pay.unmerge_cells(range_string=str(m))
for r in range(5, max(ws_pay.max_row+1, 50)):
    for c in range(1, 27):
        ws_pay.cell(r, c).value = None

# ----------------- 改造考勤表 (ws_att) -----------------
# 1. 解绑所有旧合并单元格
for m in list(ws_att.merged_cells.ranges):
    ws_att.unmerge_cells(range_string=str(m))

ai_col_idx = 4 + days_count
aj_col_idx = ai_col_idx + 1
ai_let = openpyxl.utils.get_column_letter(ai_col_idx)
aj_let = openpyxl.utils.get_column_letter(aj_col_idx)
end_day_let = openpyxl.utils.get_column_letter(3 + days_count)

# 2. 如果模板的列数大于当前月份所需的列数 (如28天导致列数缩减)，彻底删除多余的残余列与样式
if ws_att.max_column > aj_col_idx:
    ws_att.delete_cols(aj_col_idx + 1, ws_att.max_column - aj_col_idx + 10)

# 清理第5行及以下的数据
for r in range(5, max(ws_att.max_row+1, 70)):
    for c in range(1, aj_col_idx + 1):
        cell = ws_att.cell(r, c)
        cell.value = None
        cell.border = Border()

ws_att.merge_cells(f"A1:{aj_let}1")
ws_att.cell(1, 1, value="广州市轨道交通十八和二十二号线施工总承包项目部科研人员考勤表").font = Font(name="宋体", size=20, bold=True)
ws_att.cell(1, 1).alignment = align_center

ws_att.merge_cells("A2:L2")
ws_att.cell(2, 1, value="项目名称：地铁工程总承包模式下如何应用BIM平台提高项目管控效率的研究").font = font_body

ws_att.merge_cells("M2:AA2")
ws_att.cell(2, 13, value="课题合同（项目）编号：19-B-03").font = font_body

ws_att.merge_cells(f"AB2:{aj_let}2")
ws_att.cell(2, 28, value=f"{year}年{month}月份").font = font_body
ws_att.cell(2, 28).alignment = align_right

ws_att.merge_cells("A3:A4"); ws_att.cell(3, 1, value="序号").font = font_body; ws_att.cell(3, 1).alignment = align_center; ws_att.cell(3, 1).border = border_all
ws_att.merge_cells("B3:B4"); ws_att.cell(3, 2, value="姓名").font = font_body; ws_att.cell(3, 2).alignment = align_center; ws_att.cell(3, 2).border = border_all
ws_att.merge_cells("C3:C4"); ws_att.cell(3, 3, value="日期\n考勤").font = font_body; ws_att.cell(3, 3).alignment = align_center; ws_att.cell(3, 3).border = border_all

weekdays_zh = ["一", "二", "三", "四", "五", "六", "日"]
for idx, d in enumerate(days_range):
    col_pos = 4 + idx
    c_d = ws_att.cell(3, col_pos, value=d.day)
    c_d.font = font_body; c_d.alignment = align_center; c_d.border = border_all
    c_w = ws_att.cell(4, col_pos, value=weekdays_zh[d.weekday()])
    c_w.font = font_body; c_w.alignment = align_center; c_w.border = border_all
    col_let = openpyxl.utils.get_column_letter(col_pos)
    ws_att.column_dimensions[col_let].width = 3.5

ws_att.merge_cells(start_row=3, start_column=ai_col_idx, end_row=4, end_column=ai_col_idx)
ws_att.cell(3, ai_col_idx, value="出勤天数").font = font_body
ws_att.cell(3, ai_col_idx).alignment = align_center
ws_att.cell(3, ai_col_idx).border = border_all
ws_att.column_dimensions[ai_let].width = 4.58

ws_att.merge_cells(start_row=3, start_column=aj_col_idx, end_row=4, end_column=aj_col_idx)
ws_att.cell(3, aj_col_idx, value="工时率").font = font_body
ws_att.cell(3, aj_col_idx).alignment = align_center
ws_att.cell(3, aj_col_idx).border = border_all
ws_att.column_dimensions[aj_let].width = 5.42

# ----------------- 写入考勤与打点 -----------------
leader_alloc = {}
for p_idx, p in enumerate(active_people):
    nm = p['name']
    plan = assigned_plans[nm]
    r_att = 5 + 2 * p_idx
    r_res = 6 + 2 * p_idx

    ws_att.row_dimensions[r_att].height = 19.0
    ws_att.row_dimensions[r_res].height = 19.0

    ws_att.merge_cells(f"A{r_att}:A{r_res}")
    ws_att.cell(r_att, 1, value=p_idx+1).font = font_body
    ws_att.cell(r_att, 1).alignment = align_center
    ws_att.cell(r_att, 1).border = border_all
    ws_att.cell(r_res, 1).border = border_all

    ws_att.merge_cells(f"B{r_att}:B{r_res}")
    ws_att.cell(r_att, 2, value=nm).font = font_body
    ws_att.cell(r_att, 2).alignment = align_center
    ws_att.cell(r_att, 2).border = border_all
    ws_att.cell(r_res, 2).border = border_all

    ws_att.cell(r_att, 3, value="出勤").font = font_body; ws_att.cell(r_att, 3).alignment = align_center; ws_att.cell(r_att, 3).border = border_all
    ws_att.cell(r_res, 3, value="科研").font = font_body; ws_att.cell(r_res, 3).alignment = align_center; ws_att.cell(r_res, 3).border = border_all

    for c_idx in range(days_count):
        col_pos = 4 + c_idx
        c_cell = ws_att.cell(r_att, col_pos, value=marks[c_idx])
        c_cell.font = font_body; c_cell.alignment = align_center; c_cell.border = border_all
        ws_att.cell(r_res, col_pos).border = border_all

    c_tot_days = ws_att.cell(r_att, ai_col_idx, value=f"=COUNT(D{r_att}:{end_day_let}{r_att})")
    c_tot_days.font = font_body; c_tot_days.alignment = align_center; c_tot_days.border = border_all

    # 打点工时：允许标准选项 8, 4, 2, 1, 0.5
    h_target = plan['h_hours']
    
    # 采用贪心二进制分解精确匹配 h_target
    def decompose_hours(target_h):
        res = []
        rem = round(target_h * 2) / 2
        while rem >= 8:
            res.append(8); rem -= 8
        while rem >= 4:
            res.append(4); rem -= 4
        while rem >= 2:
            res.append(2); rem -= 2
        while rem >= 1:
            res.append(1); rem -= 1
        while rem >= 0.5:
            res.append(0.5); rem -= 0.5
        return res

    hours_chunks = decompose_hours(h_target)
    alloc_dict = {}
    w_indices_p = workday_indices[(p_idx * 2) % len(workday_indices):] + workday_indices[:(p_idx * 2) % len(workday_indices)]
    for chunk_idx, val in enumerate(hours_chunks):
        if chunk_idx < len(w_indices_p):
            alloc_dict[w_indices_p[chunk_idx]] = val
        else:
            alloc_dict[w_indices_p[0]] += val

    for c_idx in range(days_count):
        if c_idx in alloc_dict:
            val = alloc_dict[c_idx]
            c_h = ws_att.cell(r_res, 4 + c_idx, value=int(val) if val == int(val) else val)
            c_h.font = font_body; c_h.alignment = align_center

    c_res_days = ws_att.cell(r_res, ai_col_idx, value=f"=SUM(D{r_res}:{end_day_let}{r_res})/8")
    c_res_days.font = font_body; c_res_days.alignment = align_center; c_res_days.border = border_all
    c_res_days.number_format = '0.00_ '

    c_rate = ws_att.cell(r_res, aj_col_idx, value=f"={ai_let}{r_res}/{ai_let}{r_att}")
    c_rate.font = font_body; c_rate.alignment = align_center; c_rate.border = border_all
    c_rate.number_format = '0.0000_ '

# ----------------- 写入分摊表数据与公式 -----------------
for p_idx, p in enumerate(active_people):
    nm = p['name']
    r_pay = 5 + p_idx
    r_att_days = 5 + 2 * p_idx
    r_res_rate = 6 + 2 * p_idx

    ws_pay.row_dimensions[r_pay].height = 24.0

    c_id = ws_pay.cell(r_pay, 1, value=p_idx+1)
    c_id.font = font_body; c_id.alignment = align_center; c_id.border = border_all

    c_nm = ws_pay.cell(r_pay, 2, value=nm)
    c_nm.font = font_body; c_nm.alignment = align_center; c_nm.border = border_all

    c_d = ws_pay.cell(r_pay, 3, value=f"='201903月考勤表'!{ai_let}{r_att_days}")
    c_d.font = font_body; c_d.alignment = align_center; c_d.border = border_all

    c_r = ws_pay.cell(r_pay, 4, value=f"='201903月考勤表'!{aj_let}{r_res_rate}")
    c_r.font = font_body; c_r.alignment = align_right; c_r.border = border_all
    c_r.number_format = '0.0000'

    w_base = wages[nm]['应发工资']
    sb_item = ss_dict.get(nm, {})
    hf_base = hf_dict.get(nm, round(w_base * 0.12, 2))

    def sc(col, val, fmt='#,##0.00'):
        c = ws_pay.cell(r_pay, col, value=val)
        c.font = font_body; c.alignment = align_right; c.border = border_all
        c.number_format = fmt
        return c

    sc(5, w_base)                                    # E: 1.应发工资
    sc(6, sb_item.get('单位合计', 0.0))               # F: 2.社会保险费
    sc(7, sb_item.get('养老', 0.0))                   # G: 2.1基本养老
    sc(8, sb_item.get('工伤', 0.0))                   # H: 2.2工伤
    sc(9, sb_item.get('失业', 0.0))                   # I: 2.3失业
    sc(10, sb_item.get('医疗', 0.0))                  # J: 2.4基本医疗
    sc(11, sb_item.get('大病', 0.0))                  # K: 2.5大病医疗
    sc(12, sb_item.get('生育', 0.0))                  # L: 2.6生育保险
    sc(13, hf_base)                                  # M: 3.住房公积金
    sc(14, "")                                       # N: 4.企业年金
    sc(15, f"=SUM(E{r_pay},F{r_pay},M{r_pay},N{r_pay})") # O: 小计

    # 计入课题研发费用公式 (尾差注入研发工资，使小计完全平齐)
    comp = assigned_plans[nm].get('comp', 0.0)
    if abs(comp) > 0.00001:
        comp_str = f"+{comp:.2f}" if comp > 0 else f"{comp:.2f}"
        f_wage = f"=E{r_pay}*$D{r_pay}{comp_str}"
    else:
        f_wage = f"=E{r_pay}*$D{r_pay}"

    sc(16, f_wage)                                   # P: 1.研发工资
    sc(17, f"=F{r_pay}*$D{r_pay}")                   # Q: 2.社会保险费
    sc(18, f"=G{r_pay}*$D{r_pay}")                   # R: 2.1基本养老
    sc(19, f"=H{r_pay}*$D{r_pay}")                   # S: 2.2工伤
    sc(20, f"=I{r_pay}*$D{r_pay}")                   # T: 2.3失业
    sc(21, f"=J{r_pay}*$D{r_pay}")                   # U: 2.4基本医疗
    sc(22, f"=K{r_pay}*$D{r_pay}")                   # V: 2.5大病医疗
    sc(23, f"=L{r_pay}*$D{r_pay}")                   # W: 2.6生育保险
    sc(24, f"=M{r_pay}*$D{r_pay}")                   # X: 3.住房公积金
    sc(25, f'=IF(N{r_pay}="","",N{r_pay}*$D{r_pay})')# Y: 4.企业年金
    sc(26, f"=SUM(P{r_pay},Q{r_pay},X{r_pay},Y{r_pay})") # Z: 小计 (研发人员费用合计)

# 合计行
r_tot = 5 + len(active_people)
ws_pay.row_dimensions[r_tot].height = 24.0
ws_pay.merge_cells(f"A{r_tot}:D{r_tot}")
c_t = ws_pay.cell(r_tot, 1, value="合计")
c_t.font = font_bold; c_t.alignment = align_center
for c in range(1, 5): ws_pay.cell(r_tot, c).border = border_all

for col in range(5, 27):
    col_let = openpyxl.utils.get_column_letter(col)
    if col in [14, 25]:
        c = ws_pay.cell(r_tot, col, value=f'=IF(COUNT({col_let}5:{col_let}{r_tot-1})=0,"",SUM({col_let}5:{col_let}{r_tot-1}))')
    else:
        c = ws_pay.cell(r_tot, col, value=f"=SUM({col_let}5:{col_let}{r_tot-1})")
    c.font = font_bold; c.alignment = align_right; c.border = border_all
    c.number_format = '#,##0.00'

ws_pay.print_area = f"'201903月'!$A$1:$Z${r_tot+2}"
ws_pay.page_setup.orientation = 'landscape'
ws_pay.page_setup.paperSize = 9

r_att_last = 4 + 2 * len(active_people)
ws_att.print_area = f"'201903月考勤表'!$A$1:${aj_let}${r_att_last}"
ws_att.page_setup.orientation = 'landscape'
ws_att.page_setup.paperSize = 9

wb_out.save(out_file)
print(f"\n[OK] 成功生成 19-B-03 首期样板成果文件: {out_file}")
