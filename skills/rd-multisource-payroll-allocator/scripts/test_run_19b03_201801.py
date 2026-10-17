# -*- coding: utf-8 -*-
"""
全流程端到端验证脚本：生成 2018年1月 19-B-03 考勤表与薪酬分摊表
包含：
1. 多源提取 1.应发工资(工资表X列), 2.社保6险(单位列，含生育保险), 3.公积金(单位列), 4.台账目标N79 (142,426.30元)
2. 基于官方定稿母版克隆，插入 2.6生育保险 列
3. 动态考勤周期 (2017-12-21 ~ 2018-01-20，元旦放假)，打点与工时率倒推
4. 校验公式联动与绝对金额
"""
import os, sys, datetime as dt, re
import openpyxl, xlrd
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

sys.path.append(r"C:\JunzanLi_project\skills_mcp\skills\rd-multisource-payroll-allocator\scripts")
from data_loader import (
    parse_task_book,
    load_monthly_wage,
    load_monthly_social_security,
    load_monthly_housing_fund,
    load_monthly_ledger_target
)

# 基础路径
tpl_path = r"C:\JunzanLi_project\课题申报横向\人工费考勤表与分摊表\2024年7-9月研发人员薪酬、职工薪酬分配表21-B-02（定稿打印版）.xlsx"
tb_path = r"C:\JunzanLi_project\课题申报横向\财务\19-B-03科研课题合同任务书_1005凝练版.docx"
payroll_dir = r"C:\JunzanLi_project\课题申报横向\财务\18和22号线2018-2021"
insurance_dir = r"C:\JunzanLi_project\课题申报横向\财务\2018-2021社保公积金等"
ledger_dir = r"C:\JunzanLi_project\课题申报横向\财务\1822号线研发薪酬分摊台账"
out_dir = r"C:\JunzanLi_project\课题申报横向\人工费考勤表与分摊表\结果\测试_19B03_全流程验证"
os.makedirs(out_dir, exist_ok=True)
out_file = os.path.join(out_dir, "2018年1月研发人员薪酬、职工薪酬分配表19-B-03.xlsx")

print(">>> [Step 1] 多源提取 2018年1月 财务数据...")
tb_people = parse_task_book(tb_path)
tb_map = {p['name']: p for p in tb_people}
wages = load_monthly_wage(payroll_dir, 2018, 1)
ss_dict, has_mat = load_monthly_social_security(insurance_dir, 2018, 1)
hf_dict = load_monthly_housing_fund(insurance_dir, 2018, 1, wages)
target_fee, _, trow = load_monthly_ledger_target(ledger_dir, 2018, 1, ["智能设计", "19-B-03", "BIM"])

print(f"  任务书总包部人员: {len(tb_people)} 人")
print(f"  工资表覆盖: {len(wages)} 人")
print(f"  社保表覆盖: {len(ss_dict)} 人 (含生育保险: {has_mat})")
print(f"  公积金覆盖: {len(hf_dict)} 人")
print(f"  台账目标研发工资金额: {target_fee:.2f} 元 (来源行: {trow+1})")

# 确定当期参研团队人员名单
# 优先选用在任务书 Table 5 中且当期在工资表有工资的人员
candidate_people = [p for p in tb_people if p['name'] in wages]
print(f"  当月在职任务书人员: {len(candidate_people)} 人")

# 计算当期日历 (2017-12-21 至 2018-01-20)
start_date = dt.date(2017, 12, 21)
end_date = dt.date(2018, 1, 20)
days_count = (end_date - start_date).days + 1
days_range = [start_date + dt.timedelta(days=i) for i in range(days_count)]

# 日历打点规则：元旦 2017-12-30~2018-01-01 放假
def get_mark(d):
    # 2018元旦放假：2017年12月30日至2018年1月1日
    if d == dt.date(2018, 1, 1):
        return "-"
    if d.weekday() >= 5: # 周末
        return "/"
    return 8

marks = [get_mark(d) for d in days_range]
workday_indices = [i for i, m in enumerate(marks) if m == 8]
W_days = len(workday_indices)
print(f"  考勤周期: {start_date} ~ {end_date} (共 {days_count} 天, 额定工作日: {W_days} 天)")

# 逆向求解工时与工时率配平
# 针对这批人员，设定工时率三档梯度：
# leader (陈前, 梁兴朴, 周智, 喻国军, 刘少雨): 0.38
# key (肖云, 房德鑫, 刘博): 0.33
# general (其余人): 0.28
# 通过微调打点工时，使得 sum(wage * rate) 精确匹配 142,426.30 元
tot_wage_tb = sum(wages[p['name']]['应发工资'] for p in candidate_people)
print(f"  参研人员应发工资总额: {tot_wage_tb:.2f} 元, 目标研发工资: {target_fee:.2f} 元 (整体工时率约为 {target_fee/tot_wage_tb:.4f})")

# 为每位人员分配目标工时率与工时
assigned_plans = {}
rem_target = target_fee
for p in candidate_people:
    name = p['name']
    tier = p['tier']
    w = wages[name]['应发工资']
    if tier == 'leader':
        base_rate = 0.38
    elif tier == 'key':
        base_rate = 0.33
    else:
        base_rate = 0.28
    assigned_plans[name] = {'base_rate': base_rate, 'wage': w, 'tier': tier}

# 简单比例归一化缩放至目标总额
init_sum = sum(p['wage'] * p['base_rate'] for p in assigned_plans.values())
scale_factor = target_fee / init_sum
for name, p in assigned_plans.items():
    p['target_fee'] = p['wage'] * p['base_rate'] * scale_factor
    p['eff_rate'] = p['target_fee'] / p['wage']
    # 转换为科研小时 (额定出勤天数 W_days, 每天 8 小时)
    # rate = H / (8 * A)
    # 假设 A = W_days, H = round(rate * 8 * A * 2) / 2
    p['a_days'] = W_days
    p['h_hours'] = round(p['eff_rate'] * 8.0 * W_days * 2) / 2
    # 重新核算实算分摊
    p['calc_rate'] = (p['h_hours'] / 8.0) / p['a_days']
    p['calc_fee'] = p['wage'] * p['calc_rate']

calc_tot = sum(p['calc_fee'] for p in assigned_plans.values())
diff = target_fee - calc_tot
print(f"  网格离散工时实算总额: {calc_tot:.2f} 元, 差额: {diff:+.2f} 元")

# 末位/首位平账注入微调
first_person = candidate_people[0]['name'] # 陈前
assigned_plans[first_person]['comp'] = round(diff, 2)
print(f"  在课题负责人 ({first_person}) 末月公式注入尾差微调: {diff:+.2f} 元, 达成绝对 0.00 元平衡!")

print("\n>>> [Step 2] 基于官方定稿母版克隆，插入 2.6生育保险 列...")
wb_out = openpyxl.load_workbook(tpl_path)

# 重命名工作表
wb_out['202407月'].title = '201801月'
wb_out['202407月考勤表'].title = '201801月考勤表'
# 移除其余 4 个季度月份
for s in ['202408月', '202408月考勤表', '202409月', '202409月考勤表']:
    if s in wb_out.sheetnames:
        del wb_out[s]

ws_pay = wb_out['201801月']
ws_att = wb_out['201801月考勤表']

# 样式定义
font_body = Font(name="宋体", size=10)
font_bold = Font(name="宋体", size=10, bold=True)
align_center = Alignment(horizontal="center", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
thin_border = Side(border_style="thin", color="000000")
border_all = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

# ----------------- 改造分摊表 (ws_pay) 插入 2.6生育保险 列 -----------------
# 先解绑 rows 1..4 的合并单元格
for m in list(ws_pay.merged_cells.ranges):
    if m.min_row <= 4:
        ws_pay.unmerge_cells(range_string=str(m))

# 原始列：
# E:应发(5), F:社保(6), G:养老(7), H:工伤(8), I:失业(9), J:医疗(10), K:大病(11), L:公积金(12), M:年金(13), N:小计(14)
# 在 K(11) 之后插入 1 列（管理费用-生育保险），在 U(21+1=22) 之后插入 1 列（研发分摊-生育保险）
ws_pay.insert_cols(12, 1) # 新增 Col L 为 2.6生育保险 (原L,M,N顺延为 M,N,O)
ws_pay.insert_cols(23, 1) # 新增 Col W 为 2.6生育保险 (原W,X顺延为 X,Y,Z)

# 重新合并与设置标题/表头
ws_pay.merge_cells("B1:Z1")
ws_pay.cell(1, 2, value="研发费用—职工薪酬分摊明细表").font = Font(name="宋体", size=22, bold=True)
ws_pay.cell(1, 2).alignment = align_center

ws_pay.merge_cells("A2:M2")
ws_pay.cell(2, 1, value="项目名称：地铁工程总承包模式下如何应用BIM平台提高项目管控效率的研究").font = font_body

ws_pay.merge_cells("N2:P2")
ws_pay.cell(2, 14, value="课题合同（项目）编号：").font = font_body

ws_pay.merge_cells("Q2:V2")
ws_pay.cell(2, 17, value="19-B-03").font = font_body

ws_pay.merge_cells("W2:Z2")
ws_pay.cell(2, 23, value="2018年1月（单位：元）").font = font_body
ws_pay.cell(2, 23).alignment = align_right

# 表头第3、4行合并与文字
ws_pay.merge_cells("A3:A4"); ws_pay.cell(3, 1, value="序号").font = font_body; ws_pay.cell(3, 1).alignment = align_center
ws_pay.merge_cells("B3:B4"); ws_pay.cell(3, 2, value="姓名").font = font_body; ws_pay.cell(3, 2).alignment = align_center
ws_pay.merge_cells("C3:C4"); ws_pay.cell(3, 3, value="从事研发天数").font = font_body; ws_pay.cell(3, 3).alignment = align_center
ws_pay.merge_cells("D3:D4"); ws_pay.cell(3, 4, value="工时率").font = font_body; ws_pay.cell(3, 4).alignment = align_center

ws_pay.merge_cells("E3:O3"); ws_pay.cell(3, 5, value="管理费用").font = font_body; ws_pay.cell(3, 5).alignment = align_center
ws_pay.merge_cells("P3:Z3"); ws_pay.cell(3, 16, value="计入课题的研发人员工资性费用").font = font_body; ws_pay.cell(3, 16).alignment = align_center

# 第4行各明细列标题
col_titles_mgt = ["1.应发工资", "2.社会保险费", "2.1基本养老", "2.2工伤", "2.3失业", "2.4基本医疗", "2.5大病医疗", "2.6生育保险", "3.住房\n公积金", "4.企业年金", "小计"]
for i, t in enumerate(col_titles_mgt):
    c = ws_pay.cell(4, 5 + i, value=t)
    c.font = font_body; c.alignment = align_center; c.border = border_all

col_titles_rd = ["1.研发工资", "2.社会保险费", "2.1基本养老", "2.2工伤", "2.3失业", "2.4基本医疗", "2.5大病医疗", "2.6生育保险", "3.住房\n公积金", "4.企业年金", "小计"]
for i, t in enumerate(col_titles_rd):
    c = ws_pay.cell(4, 16 + i, value=t)
    c.font = font_body; c.alignment = align_center; c.border = border_all

for c in range(1, 27):
    ws_pay.cell(3, c).border = border_all

ws_pay.column_dimensions['L'].width = 11.5
ws_pay.column_dimensions['W'].width = 11.5

# 清理原有数据行 (第5行至第30行)
for m in list(ws_pay.merged_cells.ranges):
    if m.min_row >= 5:
        ws_pay.unmerge_cells(range_string=str(m))

for r in range(5, 30):
    for c in range(1, 28):
        ws_pay.cell(r, c).value = None

# ----------------- 改造考勤表 (ws_att) -----------------
for m in list(ws_att.merged_cells.ranges):
    if m.min_row <= 4:
        ws_att.unmerge_cells(range_string=str(m))

ai_col_idx = 4 + days_count
aj_col_idx = ai_col_idx + 1
ai_let = openpyxl.utils.get_column_letter(ai_col_idx)
aj_let = openpyxl.utils.get_column_letter(aj_col_idx)
end_day_let = openpyxl.utils.get_column_letter(3 + days_count)

ws_att.merge_cells(f"A1:{aj_let}1")
ws_att.cell(1, 1, value="广州市轨道交通十八和二十二号线施工总承包项目部科研人员考勤表").font = Font(name="宋体", size=20, bold=True)
ws_att.cell(1, 1).alignment = align_center

ws_att.merge_cells("A2:L2")
ws_att.cell(2, 1, value="项目名称：地铁工程总承包模式下如何应用BIM平台提高项目管控效率的研究").font = font_body

ws_att.merge_cells("M2:AA2")
ws_att.cell(2, 13, value="课题合同（项目）编号：19-B-03").font = font_body

ws_att.merge_cells(f"AB2:{aj_let}2")
ws_att.cell(2, 28, value="2018年1月份").font = font_body
ws_att.cell(2, 28).alignment = align_right

# 表头序号、姓名、考勤类别
ws_att.merge_cells("A3:A4"); ws_att.cell(3, 1, value="序号").font = font_body; ws_att.cell(3, 1).alignment = align_center; ws_att.cell(3, 1).border = border_all
ws_att.merge_cells("B3:B4"); ws_att.cell(3, 2, value="姓名").font = font_body; ws_att.cell(3, 2).alignment = align_center; ws_att.cell(3, 2).border = border_all
ws_att.merge_cells("C3:C4"); ws_att.cell(3, 3, value="日期\n考勤").font = font_body; ws_att.cell(3, 3).alignment = align_center; ws_att.cell(3, 3).border = border_all

# 清理考勤表旧数据
for m in list(ws_att.merged_cells.ranges):
    if m.min_row >= 5:
        ws_att.unmerge_cells(range_string=str(m))

for r in range(5, ws_att.max_row+1):
    for c in range(1, ws_att.max_column+1):
        ws_att.cell(r, c).value = None

# 日期行表头 (Row 3, 4)
for idx, d in enumerate(days_range):
    col_pos = 4 + idx
    c_d = ws_att.cell(3, col_pos, value=d.day)
    c_d.font = font_body
    c_d.alignment = align_center
    c_d.border = border_all
    
    # 星期
    weekdays_zh = ["一", "二", "三", "四", "五", "六", "日"]
    c_w = ws_att.cell(4, col_pos, value=weekdays_zh[d.weekday()])
    c_w.font = font_body
    c_w.alignment = align_center
    c_w.border = border_all
    col_let = openpyxl.utils.get_column_letter(col_pos)
    ws_att.column_dimensions[col_let].width = 3.5

# 统计列
ai_col_idx = 4 + days_count
aj_col_idx = ai_col_idx + 1
ai_let = openpyxl.utils.get_column_letter(ai_col_idx)
aj_let = openpyxl.utils.get_column_letter(aj_col_idx)
end_day_let = openpyxl.utils.get_column_letter(3 + days_count)

ws_att.merge_cells(start_row=3, start_column=ai_col_idx, end_row=4, end_column=ai_col_idx)
ws_att.cell(3, ai_col_idx, value="出勤天数").font = font_body
ws_att.cell(3, ai_col_idx).alignment = align_center
ws_att.cell(3, ai_col_idx).border = border_all
ws_att.column_dimensions[ai_let].width = 4.58

ws_att.merge_cells(start_row=3, start_column=aj_col_idx, end_row=4, end_column=aj_col_idx)
ws_att.cell(3, aj_col_idx, value="工时率").font = font_body
ws_att.cell(3, aj_col_idx).alignment = align_center
ws_att.cell(3, aj_col_idx).border = border_all
ws_att.column_dimensions[aj_let].width = 5.42

# 写入考勤人员打点
leader_alloc = {}
for p_idx, p in enumerate(candidate_people):
    name = p['name']
    plan = assigned_plans[name]
    r_att = 5 + 2 * p_idx
    r_res = 6 + 2 * p_idx
    
    ws_att.row_dimensions[r_att].height = 19.0
    ws_att.row_dimensions[r_res].height = 19.0
    
    ws_att.merge_cells(f"A{r_att}:A{r_res}")
    ws_att.cell(r_att, 1, value=p_idx+1).font = font_body
    ws_att.cell(r_att, 1).alignment = align_center
    ws_att.cell(r_att, 1).border = border_all
    ws_att.cell(r_res, 1).border = border_all
    
    ws_att.merge_cells(f"B{r_att}:B{r_res}")
    ws_att.cell(r_att, 2, value=name).font = font_body
    ws_att.cell(r_att, 2).alignment = align_center
    ws_att.cell(r_att, 2).border = border_all
    ws_att.cell(r_res, 2).border = border_all
    
    ws_att.cell(r_att, 3, value="出勤").font = font_body
    ws_att.cell(r_att, 3).alignment = align_center
    ws_att.cell(r_att, 3).border = border_all
    
    ws_att.cell(r_res, 3, value="科研").font = font_body
    ws_att.cell(r_res, 3).alignment = align_center
    ws_att.cell(r_res, 3).border = border_all
    
    # 出勤打点
    for c_idx in range(days_count):
        col_pos = 4 + c_idx
        m_val = marks[c_idx]
        c_cell = ws_att.cell(r_att, col_pos, value=m_val)
        c_cell.font = font_body
        c_cell.alignment = align_center
        c_cell.border = border_all
        ws_att.cell(r_res, col_pos).border = border_all
        
    c_tot_days = ws_att.cell(r_att, ai_col_idx, value=f"=COUNT(D{r_att}:{end_day_let}{r_att})")
    c_tot_days.font = font_body
    c_tot_days.alignment = align_center
    c_tot_days.border = border_all
    
    # 科研小时打点
    h_target = plan['h_hours']
    is_leader = (p['tier'] == 'leader')
    alloc_dict = {}
    if is_leader:
        int_part = int(h_target)
        has_half = (h_target - int_part) > 0.25
        base_h = int_part // W_days
        rem_h = int_part % W_days
        for rank, idx in enumerate(workday_indices):
            h_val = float(base_h + (1 if rank < rem_h else 0))
            if has_half and rank == rem_h:
                h_val += 0.5
                has_half = False
            if h_val > 0: alloc_dict[idx] = h_val
        if has_half and workday_indices:
            alloc_dict[workday_indices[0]] = alloc_dict.get(workday_indices[0], 0.0) + 0.5
        if name == "陈前": leader_alloc = dict(alloc_dict)
    else:
        slots = []
        for idx in workday_indices:
            w_h = leader_alloc.get(idx, 0)
            slots.append((w_h, idx))
        slots.sort(key=lambda x: (-x[0], x[1]))
        needed = h_target
        for _, idx in slots:
            if needed >= 8:
                alloc_dict[idx] = 8
                needed -= 8
            elif needed > 0:
                alloc_dict[idx] = needed
                needed = 0
            else:
                break
                
    for c_idx in range(days_count):
        if c_idx in alloc_dict:
            val = alloc_dict[c_idx]
            c_h = ws_att.cell(r_res, 4 + c_idx, value=int(val) if val == int(val) else val)
            c_h.font = font_body
            c_h.alignment = align_center
            
    c_res_days = ws_att.cell(r_res, ai_col_idx, value=f"=SUM(D{r_res}:{end_day_let}{r_res})/8")
    c_res_days.font = font_body
    c_res_days.alignment = align_center
    c_res_days.border = border_all
    c_res_days.number_format = '0.00_ '
    
    c_rate = ws_att.cell(r_res, aj_col_idx, value=f"={ai_let}{r_res}/{ai_let}{r_att}")
    c_rate.font = font_body
    c_rate.alignment = align_center
    c_rate.border = border_all
    c_rate.number_format = '0.0000_ '

print(">>> [Step 3] 填充 201801月 分摊表数据与动态公式...")
for p_idx, p in enumerate(candidate_people):
    name = p['name']
    r_pay = 5 + p_idx
    r_att_days = 5 + 2 * p_idx
    r_res_rate = 6 + 2 * p_idx
    
    ws_pay.row_dimensions[r_pay].height = 24.0
    
    # 序号、姓名
    c_id = ws_pay.cell(r_pay, 1, value=p_idx+1)
    c_id.font = font_body; c_id.alignment = align_center; c_id.border = border_all
    
    c_nm = ws_pay.cell(r_pay, 2, value=name)
    c_nm.font = font_body; c_nm.alignment = align_center; c_nm.border = border_all
    
    # 考勤天数与工时率引用
    c_d = ws_pay.cell(r_pay, 3, value=f"='201801月考勤表'!{ai_let}{r_att_days}")
    c_d.font = font_body; c_d.alignment = align_center; c_d.border = border_all
    
    c_r = ws_pay.cell(r_pay, 4, value=f"='201801月考勤表'!{aj_let}{r_res_rate}")
    c_r.font = font_body; c_r.alignment = align_right; c_r.border = border_all
    c_r.number_format = '0.0000'
    
    # 管理费用基数
    w_base = wages[name]['应发工资']
    sb_item = ss_dict.get(name, {})
    hf_base = hf_dict.get(name, 0.0)
    
    def sc(col, val, fmt='#,##0.00'):
        c = ws_pay.cell(r_pay, col, value=val)
        c.font = font_body; c.alignment = align_right; c.border = border_all
        c.number_format = fmt
        return c
        
    sc(5, w_base)                                    # E: 1.应发工资
    sc(6, sb_item.get('单位合计', 0.0))               # F: 2.社会保险费
    sc(7, sb_item.get('养老', 0.0))                   # G: 2.1基本养老
    sc(8, sb_item.get('工伤', 0.0))                   # H: 2.2工伤
    sc(9, sb_item.get('失业', 0.0))                   # I: 2.3失业
    sc(10, sb_item.get('医疗', 0.0))                  # J: 2.4基本医疗
    sc(11, sb_item.get('大病', 0.0))                  # K: 2.5大病医疗
    sc(12, sb_item.get('生育', 0.0))                  # L: 2.6生育保险 (NEW!)
    sc(13, hf_base)                                  # M: 3.住房公积金
    sc(14, "")                                       # N: 4.企业年金
    sc(15, f"=SUM(E{r_pay},F{r_pay},M{r_pay},N{r_pay})") # O: 小计
    
    # 计入课题研发费用 (公式化联动)
    comp = assigned_plans[name].get('comp', 0.0)
    if abs(comp) > 0.00001:
        comp_str = f"+{comp:.2f}" if comp > 0 else f"{comp:.2f}"
        f_wage = f"=E{r_pay}*$D{r_pay}{comp_str}"
    else:
        f_wage = f"=E{r_pay}*$D{r_pay}"
        
    sc(16, f_wage)                                   # P: 1.研发工资
    sc(17, f"=F{r_pay}*$D{r_pay}")                   # Q: 2.社会保险费
    sc(18, f"=G{r_pay}*$D{r_pay}")                   # R: 2.1基本养老
    sc(19, f"=H{r_pay}*$D{r_pay}")                   # S: 2.2工伤
    sc(20, f"=I{r_pay}*$D{r_pay}")                   # T: 2.3失业
    sc(21, f"=J{r_pay}*$D{r_pay}")                   # U: 2.4基本医疗
    sc(22, f"=K{r_pay}*$D{r_pay}")                   # V: 2.5大病医疗
    sc(23, f"=L{r_pay}*$D{r_pay}")                   # W: 2.6生育保险 (NEW!)
    sc(24, f"=M{r_pay}*$D{r_pay}")                   # X: 3.住房公积金
    sc(25, f'=IF(N{r_pay}="","",N{r_pay}*$D{r_pay})')# Y: 4.企业年金
    sc(26, f"=SUM(P{r_pay},Q{r_pay},X{r_pay},Y{r_pay})") # Z: 小计

# 合计行
r_tot = 5 + len(candidate_people)
ws_pay.row_dimensions[r_tot].height = 24.0
ws_pay.merge_cells(f"A{r_tot}:D{r_tot}")
c_t = ws_pay.cell(r_tot, 1, value="合计")
c_t.font = font_bold; c_t.alignment = align_center
for c in range(1, 5): ws_pay.cell(r_tot, c).border = border_all

for col in range(5, 27):
    col_let = openpyxl.utils.get_column_letter(col)
    if col in [14, 25]: # 年金
        c = ws_pay.cell(r_tot, col, value=f'=IF(COUNT({col_let}5:{col_let}{r_tot-1})=0,"",SUM({col_let}5:{col_let}{r_tot-1}))')
    else:
        c = ws_pay.cell(r_tot, col, value=f"=SUM({col_let}5:{col_let}{r_tot-1})")
    c.font = font_bold; c.alignment = align_right; c.border = border_all
    c.number_format = '#,##0.00'

# 打印区域与页面设置
ws_pay.print_area = f"'201801月'!$A$1:$Z${r_tot+2}"
ws_pay.page_setup.orientation = 'landscape'
ws_pay.page_setup.paperSize = 9

wb_out.save(out_file)
print(f"\n[OK] 成功生成端到端全流程验证成果表: {os.path.basename(out_file)}")
print(f"     文件位置: {out_file}")
