# -*- coding: utf-8 -*-
import openpyxl, sys

f = r"C:\JunzanLi_project\课题申报横向\人工费考勤表与分摊表\结果\测试_19B03_全流程验证\2018年1月研发人员薪酬、职工薪酬分配表19-B-03.xlsx"
wb = openpyxl.load_workbook(f, data_only=False)
ws_p = wb['201801月']
ws_a = wb['201801月考勤表']

print("=== 201801月 考勤表验证 ===")
print("Title:", ws_a.cell(1, 1).value)
print("Subheader:", ws_a.cell(2, 1).value)
print("Project Code:", ws_a.cell(2, 13).value)
print("Date:", ws_a.cell(2, 28).value)

att_results = {}
for p_idx in range(18):
    r_att = 5 + 2 * p_idx
    r_res = 6 + 2 * p_idx
    name = ws_a.cell(r_att, 2).value
    days = sum(1 for c in range(4, 35) if ws_a.cell(r_att, c).value == 8)
    hours = sum(float(ws_a.cell(r_res, c).value or 0) for c in range(4, 35))
    rate = (hours / 8.0) / days if days > 0 else 0.0
    att_results[name] = {'days': days, 'hours': hours, 'rate': rate}
    rate_formula = ws_a.cell(r_res, 36).value
    if p_idx < 5 or p_idx >= 15:
        print(f"  {name:6s} | 出勤天数={days} | 科研小时={hours:4.1f} | 算得工时率={rate:.4f} (公式={rate_formula})")

print("\n=== 201801月 分摊表核算与平账核验 ===")
print("Title:", ws_p.cell(1, 2).value)
print("Header L4 (管理-生育):", ws_p.cell(4, 12).value)
print("Header W4 (研发-生育):", ws_p.cell(4, 23).value)

tot_wage = 0.0
tot_rd_wage = 0.0
tot_sb = 0.0
tot_rd_sb = 0.0
tot_mat = 0.0
tot_rd_mat = 0.0
tot_gjj = 0.0
tot_rd_gjj = 0.0
tot_rd_all = 0.0

for r in range(5, 23):
    name = ws_p.cell(r, 2).value
    rate_val = att_results[name]['rate']
    
    w = float(ws_p.cell(r, 5).value or 0)
    sb = float(ws_p.cell(r, 6).value or 0)
    mat = float(ws_p.cell(r, 12).value or 0)
    gjj = float(ws_p.cell(r, 13).value or 0)
    
    f_rd_w = str(ws_p.cell(r, 16).value)
    comp = 0.0
    if '+' in f_rd_w: comp = float(f_rd_w.split('+')[1])
    elif '-' in f_rd_w and '*$D' in f_rd_w: comp = -float(f_rd_w.split('-')[-1])
    
    rd_w = w * rate_val + comp
    rd_sb = sb * rate_val
    rd_mat = mat * rate_val
    rd_gjj = gjj * rate_val
    rd_all = rd_w + rd_sb + rd_gjj
    
    tot_wage += w
    tot_rd_wage += rd_w
    tot_sb += sb
    tot_rd_sb += rd_sb
    tot_mat += mat
    tot_rd_mat += rd_mat
    tot_gjj += gjj
    tot_rd_gjj += rd_gjj
    tot_rd_all += rd_all
    
    if r < 10 or r >= 20:
        print(f"  {name:6s} | 应发={w:8.2f} | 研发工资={rd_w:8.2f} | 研发社保={rd_sb:7.2f} (含生育={rd_mat:6.2f}) | 研发公积金={rd_gjj:7.2f} | 研发合计={rd_all:8.2f}")

target_n79 = 142426.30
diff = tot_rd_wage - target_n79
status = "100% 绝对平衡 (0.00)" if abs(diff) < 0.005 else "存在误差"

print("=" * 85)
print(f"汇总合计应发工资: {tot_wage:.2f} 元")
print(f"实际计算研发工资总额: {tot_rd_wage:.2f} 元")
print(f"台账法定目标金额 N79: {target_n79:.2f} 元")
print(f"绝对平衡差额: {diff:+.4f} 元 (状态: {status})")
print(f"研发人员社保合计: {tot_rd_sb:.2f} 元 (其中 2.6生育保险: {tot_rd_mat:.2f} 元)")
print(f"研发人员公积金合计: {tot_rd_gjj:.2f} 元")
print(f"计入课题全要素总额 (工资+社保+公积金): {tot_rd_all:.2f} 元")
print("=" * 85)
